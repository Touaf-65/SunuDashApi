"""
Lecture des fichiers et préparation des données (lot I1).

Chaque fichier est relu tel quel (colonnes d'origine, pour le rapport) et traduit vers les noms communs
de columns.py. Les deux tableaux partagent le même index : une ligne préparée renvoie toujours
à sa ligne d'origine (fichier, feuille, numéro de ligne Excel).
"""
import os
from dataclasses import dataclass, field

import pandas as pd

from .columns import STAT, RECAP, KIND_LABELS, match_columns, missing_columns, detect_kind
from .normalize import clean_text, text_key, claim_key, parse_amounts, parse_dates

# Colonnes ajoutées aux lignes d'origine pour les retrouver dans le fichier
ORIGIN_COLUMNS = ['Fichier', 'Feuille', 'Ligne']

# Premières lignes d'un tableur : en-tête en ligne 1, données à partir de la ligne 2
FIRST_DATA_ROW = 2


class SourceError(Exception):
    """Fichier illisible, de mauvaise nature ou incomplet. Le message est destiné à l'utilisateur."""


def _ext(filename):
    return os.path.splitext(filename or '')[1].lower()


def is_lock_file(filename):
    """Fichier verrou d'Excel (~$classeur.xlsx, .~classeur.xlsx) : jamais un vrai fichier de données."""
    base = os.path.basename(filename or '')
    return base.startswith('~$') or base.startswith('.~')


def list_sheets(path, filename=None):
    """Noms des feuilles d'un classeur ([] pour un CSV, qui n'en a qu'une)."""
    ext = _ext(filename or path)
    try:
        if ext == '.xlsx':
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True)
            try:
                return list(wb.sheetnames)
            finally:
                wb.close()
        if ext == '.xls':
            import xlrd
            book = xlrd.open_workbook(path, on_demand=True)
            try:
                return list(book.sheet_names())
            finally:
                book.release_resources()
    except Exception as exc:
        raise SourceError(f"{os.path.basename(filename or path)} : classeur illisible ({exc}).")
    return []


def _read_csv(path):
    last_error = None
    for encoding in ('utf-8-sig', 'cp1252'):
        try:
            return pd.read_csv(path, sep=None, engine='python', dtype=str, encoding=encoding)
        except UnicodeDecodeError as exc:
            last_error = exc
    raise last_error


def read_table(path, filename=None, sheet=None):
    """Tableau brut d'un fichier (une feuille), avec les colonnes Fichier / Feuille / Ligne en tête."""
    filename = filename or os.path.basename(path)
    ext = _ext(filename)
    try:
        if ext in ('.xlsx', '.xls'):
            df = pd.read_excel(path, sheet_name=sheet if sheet is not None else 0)
        elif ext == '.csv':
            df = _read_csv(path)
        else:
            raise SourceError(f"{filename} : format non accepté (.xlsx, .xls ou .csv).")
    except SourceError:
        raise
    except ValueError as exc:
        if sheet is not None and 'not found' in str(exc).lower():
            raise SourceError(f"{filename} : la feuille « {sheet} » n'existe pas.")
        raise SourceError(f"{filename} : fichier illisible ({exc}).")
    except Exception as exc:
        raise SourceError(f"{filename} : fichier illisible ({exc}).")

    df.columns = [str(c) for c in df.columns]
    df.insert(0, 'Ligne', range(FIRST_DATA_ROW, FIRST_DATA_ROW + len(df)))
    df.insert(0, 'Feuille', sheet or '')
    df.insert(0, 'Fichier', filename)
    data_cols = [c for c in df.columns if c not in ORIGIN_COLUMNS]
    return df.dropna(how='all', subset=data_cols)


def check_kind(raw, expected, filename):
    """Vérifie qu'un fichier est bien de la nature attendue et complet, sinon SourceError."""
    columns = [c for c in raw.columns if c not in ORIGIN_COLUMNS]
    found = detect_kind(columns)
    label = KIND_LABELS[expected]
    if found is None:
        raise SourceError(f"{filename} n'est pas un {label} : colonnes non reconnues.")
    if found != expected:
        raise SourceError(f"{filename} a été déposé comme {label}, mais c'est un {KIND_LABELS[found]}.")
    missing = missing_columns(columns, expected)
    if missing:
        raise SourceError(f"{filename} : colonne(s) obligatoire(s) absente(s) ou sans titre : {', '.join(missing)}.")
    if raw.empty:
        raise SourceError(f"{filename} : aucune ligne de données.")


def _canonical(raw, kind):
    mapping = match_columns([c for c in raw.columns if c not in ORIGIN_COLUMNS], kind)
    canon = pd.DataFrame(index=raw.index)
    for name, col in mapping.items():
        canon[name] = raw[col]
    return canon, mapping


@dataclass
class Prepared:
    """Données d'un côté (statistique ou récap), prêtes pour le rapprochement."""
    raw: pd.DataFrame                   # lignes d'origine (colonnes d'origine + Fichier / Feuille / Ligne)
    data: pd.DataFrame                  # mêmes lignes, noms communs et valeurs normalisées
    mapping: dict                       # nom commun -> colonne d'origine
    anomalies: pd.DataFrame = field(default_factory=pd.DataFrame)  # index, motif
    duplicate_rows_removed: int = 0     # récap : lignes identiques retirées (pages qui se chevauchent)
    files: list = field(default_factory=list)


def _anomaly(index, reason):
    return pd.DataFrame({'motif': reason}, index=index)


def prepare_stat(raw):
    """Statistique : une ligne par acte. Les lignes strictement identiques sont gardées (actes répétés légitimes)."""
    data, mapping = _canonical(raw, STAT)
    data['claim_id'] = claim_key(data['claim_id'])
    data['settlement_date'] = parse_dates(data['settlement_date'])
    data['incident_date'] = parse_dates(data['incident_date'])
    data['claimed_amount'] = parse_amounts(data['claimed_amount'])
    data['reimbursed_amount'] = parse_amounts(data['reimbursed_amount'])
    data['claim_status_key'] = data['claim_status'].map(text_key)
    data['insured_status'] = data['insured_status'].map(lambda v: (clean_text(v) or '').upper())

    problems = []
    for mask, reason in (
        (data['claim_id'].isna(), 'Numéro de sinistre absent'),
        (data['settlement_date'].isna(), 'Date de règlement absente ou illisible'),
        (data['claimed_amount'].isna(), 'Montant facturé absent ou illisible'),
        (data['reimbursed_amount'].isna(), 'Montant remboursé absent ou illisible'),
    ):
        if mask.any():
            problems.append(_anomaly(data.index[mask], reason))
    anomalies = _merge_anomalies(problems)
    return Prepared(raw=raw, data=data, mapping=mapping, anomalies=anomalies,
                    files=list(dict.fromkeys(raw['Fichier'])))


def prepare_recaps(raws):
    """Récaps : une ligne par sinistre. Plusieurs fichiers (pages d'un même export) sont fusionnés,
    puis les lignes strictement identiques retirées."""
    raw = pd.concat(raws, ignore_index=True)
    files = list(dict.fromkeys(raw['Fichier']))
    data, mapping = _canonical(raw, RECAP)
    before = len(raw)
    keep = ~raw.duplicated(subset=list(mapping.values()), keep='first')
    raw, data = raw[keep], data[keep]
    removed = before - len(raw)

    raw_dates = data['settlement_date']
    data['claim_id'] = claim_key(data['claim_id'])
    data['settlement_date'] = parse_dates(raw_dates)
    data['claimed_amount'] = parse_amounts(data['claimed_amount'])
    data['reimbursed_amount'] = parse_amounts(data['reimbursed_amount'])
    # Non payé : pas de date de règlement (« - » ou vide) -> exclu du rapprochement, listé dans le rapport
    data['unpaid'] = raw_dates.map(clean_text).isna()

    paid = ~data['unpaid']
    problems = []
    for mask, reason in (
        (data['claim_id'].isna(), 'Numéro de sinistre absent'),
        (paid & data['settlement_date'].isna(), 'Date de règlement illisible'),
        (paid & data['claimed_amount'].isna(), 'Montant réclamé absent ou illisible'),
        (paid & data['reimbursed_amount'].isna(), 'Montant remboursé absent ou illisible'),
    ):
        if mask.any():
            problems.append(_anomaly(data.index[mask], reason))
    return Prepared(raw=raw, data=data, mapping=mapping, anomalies=_merge_anomalies(problems),
                    duplicate_rows_removed=removed, files=files)


def _merge_anomalies(frames):
    if not frames:
        return pd.DataFrame({'motif': pd.Series(dtype=str)})
    merged = pd.concat(frames)
    return merged.groupby(level=0)['motif'].agg(' ; '.join).to_frame()


def load_stat(path, filename=None, sheet=None):
    filename = filename or os.path.basename(path)
    raw = read_table(path, filename, sheet)
    check_kind(raw, STAT, f"{filename}" + (f" (feuille « {sheet} »)" if sheet else ''))
    return prepare_stat(raw)


def load_recaps(files):
    """`files` : liste de (chemin, nom affiché). Les fichiers verrous d'Excel sont ignorés."""
    raws = []
    for path, filename in files:
        filename = filename or os.path.basename(path)
        if is_lock_file(filename):
            continue
        raw = read_table(path, filename)
        check_kind(raw, RECAP, filename)
        raws.append(raw)
    if not raws:
        raise SourceError("Aucun fichier récap exploitable.")
    return prepare_recaps(raws)
